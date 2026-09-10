# -*- coding: utf-8 -*-
"""王福鑫 · 环境设计作品集 PPT — 仿站黑白橙风"""

from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import copy

OUT = Path(r"E:\作品集\王福鑫-环境设计作品集.pptx")
ASSETS = Path(r"E:\作品集\00000\仿mimo作品集官网\assets")
TMP = Path(r"E:\作品集\00000\仿mimo作品集官网\assets\_ppt_tmp")
TMP.mkdir(parents=True, exist_ok=True)

# 16:9
SW, SH = Inches(13.333), Inches(7.5)
INK = RGBColor(0x11, 0x11, 0x13)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
FAINT = RGBColor(0x9A, 0x9A, 0x9A)
LINE = RGBColor(0xE8, 0xE8, 0xE8)
ORANGE = RGBColor(0xFF, 0x69, 0x00)
BG = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xF7, 0xF7, 0xF8)

FONT = "Microsoft YaHei"
FONT_EN = "Segoe UI"


def set_run_font(run, name=FONT, size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return shape


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    return sh


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf.paragraphs[0].alignment = align
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, font, size, bold, color)
    # vertical center optional via tf
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    if bodyPr is not None:
        anchor_map = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}
        bodyPr.set("anchor", anchor_map.get(anchor, "t"))
    return box


def prep_image(src: Path, max_w=1400) -> Path | None:
    if not src or not src.exists():
        return None
    dest = TMP / (src.stem + "_ppt.jpg")
    try:
        im = Image.open(src)
        if im.mode in ("RGBA", "P"):
            im = im.convert("RGB")
        w, h = im.size
        if w > max_w:
            im = im.resize((max_w, int(h * max_w / w)), Image.Resampling.LANCZOS)
        im.save(dest, "JPEG", quality=88, optimize=True)
        return dest
    except Exception as e:
        print("img fail", src, e)
        return None


def add_cover_image(slide, path: Path, x, y, w, h):
    p = prep_image(path)
    if not p:
        add_rect(slide, x, y, w, h, SOFT, LINE)
        return
    # cover-fit via crop
    im = Image.open(p)
    iw, ih = im.size
    target = w / h
    src = iw / ih
    pic = slide.shapes.add_picture(str(p), x, y, width=w, height=h)
    if src > target:
        # too wide - crop sides
        crop = (1 - target / src) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    elif src < target:
        crop = (1 - src / target) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop


def footer(slide, page=None):
    add_rect(slide, Inches(0.55), Inches(7.15), Inches(12.2), Emu(9525), LINE, None)
    add_text(slide, Inches(0.55), Inches(7.18), Inches(6), Inches(0.28), "王福鑫 · 环境设计作品集", 9, False, FAINT)
    if page is not None:
        add_text(slide, Inches(10.5), Inches(7.18), Inches(2.3), Inches(0.28), f"{page:02d}", 9, False, FAINT, PP_ALIGN.RIGHT, FONT_EN)


def orange_bar(slide, x, y, w=Inches(0.55)):
    add_rect(slide, x, y, w, Emu(28575), ORANGE, None)


def section_label(slide, x, y, zh, en):
    orange_bar(slide, x, y)
    add_text(slide, x, y + Inches(0.1), Inches(4), Inches(0.3), zh, 11, True, ORANGE)
    add_text(slide, x + Inches(1.2), y + Inches(0.1), Inches(4), Inches(0.3), en, 10, False, FAINT, font=FONT_EN)


def title_bar(slide, x, y, w, h, title, tag=None, desc=None):
    """黑框白底 keycap style card"""
    add_rect(slide, x, y, w, h, WHITE, INK, 2.25)
    add_rect(slide, x + Inches(0.06), y + Inches(0.06), w, h, None, INK, 2.25)  # offset shadow hint
    ty = y + Inches(0.28)
    if tag:
        add_text(slide, x + Inches(0.28), ty, w - Inches(0.5), Inches(0.28), tag, 10, False, FAINT, font=FONT_EN)
        ty += Inches(0.3)
    add_text(slide, x + Inches(0.28), ty, w - Inches(0.5), Inches(0.45), title, 20, True, INK)
    if desc:
        add_text(slide, x + Inches(0.28), ty + Inches(0.48), w - Inches(0.55), Inches(0.55), desc, 12, False, MUTED)


def list_row(slide, x, y, w, num, title, sub, accent=False):
    h = Inches(0.72)
    bgc = INK if accent else WHITE
    tc = WHITE if accent else INK
    sc = RGBColor(0xAA, 0xAA, 0xAA) if accent else MUTED
    add_rect(slide, x, y, w, h, bgc, INK, 2)
    add_rect(slide, x + Inches(0.04), y + Inches(0.04), w, h, None, INK, 2)
    add_text(slide, x + Inches(0.22), y + Inches(0.18), Inches(0.5), Inches(0.4), num, 14, True, ORANGE if not accent else ORANGE, font=FONT_EN)
    add_text(slide, x + Inches(0.8), y + Inches(0.12), w - Inches(1.4), Inches(0.3), title, 15, True, tc)
    add_text(slide, x + Inches(0.8), y + Inches(0.4), w - Inches(1.4), Inches(0.28), sub, 11, False, sc)
    add_text(slide, x + w - Inches(0.45), y + Inches(0.22), Inches(0.3), Inches(0.3), "→", 14, False, tc, PP_ALIGN.RIGHT)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]

    # ========== 1 Cover ==========
    s = prs.slides.add_slide(blank)
    add_bg(s)
    # left orange edge
    add_rect(s, 0, 0, Inches(0.12), SH, ORANGE, None)
    add_text(s, Inches(0.9), Inches(1.55), Inches(8), Inches(0.35), "PORTFOLIO · 2026 · QINGDAO", 12, True, ORANGE, font=FONT_EN)
    add_text(s, Inches(0.9), Inches(2.15), Inches(10), Inches(1.1), "方寸之间", 56, True, INK)
    add_text(s, Inches(1.6), Inches(3.2), Inches(10), Inches(1.1), "想象之上", 56, True, RGBColor(0x44, 0x44, 0x44))
    add_rect(s, Inches(0.92), Inches(4.5), Inches(1.1), Emu(28575), ORANGE, None)
    add_text(s, Inches(0.9), Inches(4.75), Inches(9), Inches(0.4), "你好，我是 王福鑫 · 环境设计", 18, False, MUTED)
    add_text(s, Inches(0.9), Inches(5.35), Inches(9), Inches(0.35), "景观方案 · 交互原型 · AI 影像 · 数字孪生", 13, False, FAINT)
    add_text(s, Inches(0.9), Inches(6.3), Inches(8), Inches(0.3), "青岛理工大学 · 艺术与设计学院", 12, False, FAINT)
    add_text(s, Inches(0.9), Inches(6.65), Inches(8), Inches(0.3), "空间、设计，与一切有意思的事，我都感兴趣。", 14, True, ORANGE)
    notes(s, "封面：方寸之间，想象之上。环境设计作品集。")

    # ========== 2 About ==========
    s = prs.slides.add_slide(blank)
    add_bg(s)
    section_label(s, Inches(0.7), Inches(0.55), "关于我", "ABOUT")
    add_text(s, Inches(0.7), Inches(1.1), Inches(11), Inches(0.6), "HELLO, I'M 王福鑫", 32, True, INK, font=FONT_EN)
    orange_bar(s, Inches(0.72), Inches(1.75), Inches(1.4))

    add_rect(s, Inches(0.7), Inches(2.15), Inches(7.6), Inches(4.2), SOFT, LINE)
    body = (
        "关注空间与人的关系——从文脉、生态到日常使用，习惯先分析、再推演，最后把讲得清楚的空间表达出来。\n\n"
        "主线是景观方案：城市中央公园、山海乡村更新、企业共享花园。同时把同一套空间思维延伸到 AI 影像、可交互网站、室内 3D 与数字孪生——设计不止一种媒介。\n\n"
        "也做能点开就用的小工具：像素画板、全屋工坊，以及上线的展示站。欢迎交流方案、原型与合作。"
    )
    add_text(s, Inches(1.0), Inches(2.4), Inches(7.1), Inches(3.5), body, 14, False, RGBColor(0x33, 0x33, 0x33))
    add_text(s, Inches(1.0), Inches(5.7), Inches(7), Inches(0.4), "— 王福鑫 · WANG FUXIN", 13, True, INK)

    add_rect(s, Inches(8.6), Inches(2.15), Inches(4.0), Inches(4.2), WHITE, LINE)
    meta = [("院校", "青岛理工大学 · 艺术与设计学院"), ("专业", "环境设计 · 环设 233"),
            ("方向", "景观 · 室内 · AI 影像 · 交互 · 孪生"), ("状态", "在读")]
    my = Inches(2.4)
    for k, v in meta:
        add_text(s, Inches(8.9), my, Inches(1.2), Inches(0.28), k, 11, False, FAINT)
        add_text(s, Inches(8.9), my + Inches(0.28), Inches(3.5), Inches(0.4), v, 13, True, INK)
        add_rect(s, Inches(8.9), my + Inches(0.72), Inches(3.4), Emu(9525), LINE, None)
        my += Inches(0.9)
    footer(s, 2)
    notes(s, "关于我：景观主线 + 多媒介实验。")

    # ========== 3 Selected works overview ==========
    s = prs.slides.add_slide(blank)
    add_bg(s)
    section_label(s, Inches(0.7), Inches(0.5), "精选作品", "SELECTED WORKS")
    add_text(s, Inches(0.7), Inches(1.0), Inches(11), Inches(0.55), "三大板块", 28, True, INK)
    add_text(s, Inches(0.7), Inches(1.65), Inches(11), Inches(0.3), "景观设计 · 交互数字孪生 · AI 视频影像", 13, False, MUTED)

    cols = [
        (Inches(0.7), "01", "景观设计", "LANDSCAPE", "青峦望海 / 流筑未来", ASSETS / "flow-garden.png"),
        (Inches(4.85), "02", "交互数字孪生", "INTERACTIVE / TWIN", "智慧社区 · 全屋工坊", ASSETS / "twin-2.png"),
        (Inches(9.0), "03", "AI 视频影像", "AI SHOWREEL", "生成影像实验场", ASSETS / "ai-1.png"),
    ]
    for x, no, zh, en, sub, img in cols:
        add_text(s, x, Inches(2.15), Inches(0.6), Inches(0.3), no, 12, True, ORANGE, font=FONT_EN)
        add_text(s, x + Inches(0.5), Inches(2.12), Inches(3), Inches(0.35), zh, 14, True, INK)
        add_text(s, x, Inches(2.5), Inches(3.5), Inches(0.25), en, 9, False, FAINT, font=FONT_EN)
        add_cover_image(s, img, x, Inches(2.9), Inches(3.4), Inches(3.5))
        add_rect(s, x, Inches(2.9), Inches(3.4), Inches(3.5), None, INK, 2)
    footer(s, 3)
    notes(s, "精选作品总览：三栏分类。")

    # ========== 4-6 Landscape projects ==========
    projects = [
        {
            "no": "01", "title": "青峦望海 · 青山村", "en": "LAOSHAN · QINGSHAN VILLAGE",
            "desc": "崂山六百年渔村老村委片区低干预改造。保留红瓦石墙与海草房记忆，让山海在台地、剧场与观景平台之间重新相连。",
            "meta": "青岛 · 崂山  |  乡村景观更新  |  2026.4",
            "imgs": ["assets/hero-qingshan.jpg", "assets/render-dusk.jpg", "assets/render-entrance.jpg"],
        },
        {
            "no": "02", "title": "泉汇西枢 · 济南西中央公园", "en": "JINAN WEST CENTRAL PARK",
            "desc": "立足济南西站客流与地铁 TOD，沿中央公园中轴组织广场、水岸与林带，让换乘的人潮在一座公园里慢下来。",
            "meta": "济南 · 西站片区  |  城市中央公园  |  景观设计",
            "imgs": ["aerial.jpg", "render-1.jpg", "plan-1.jpg"],
        },
        {
            "no": "03", "title": "流·筑未来 · 企业共享花园", "en": "CORPORATE SHARED GARDEN",
            "desc": "以「流动」与「构筑」为双母题——曲线园路串联环形空间，混凝土被转译为坐凳与艺术构筑，织成企业展示、员工交流与生态休闲共栖的花园。",
            "meta": "淄博 · 产业片区  |  办公景观  |  综合实习 2026",
            "imgs": ["assets/flow-garden.png", "assets/landscape-1.png", "assets/landscape-2.png"],
        },
    ]
    heji = Path(r"E:\作品集\合集")
    for i, pr in enumerate(projects):
        s = prs.slides.add_slide(blank)
        add_bg(s)
        section_label(s, Inches(0.7), Inches(0.45), "景观设计", "LANDSCAPE")
        add_text(s, Inches(0.7), Inches(0.95), Inches(1.2), Inches(0.35), pr["no"], 14, True, ORANGE, font=FONT_EN)
        add_text(s, Inches(1.5), Inches(0.9), Inches(10), Inches(0.5), pr["title"], 26, True, INK)
        add_text(s, Inches(1.5), Inches(1.45), Inches(10), Inches(0.28), pr["en"], 10, False, FAINT, font=FONT_EN)
        add_text(s, Inches(0.7), Inches(1.95), Inches(11.5), Inches(0.7), pr["desc"], 13, False, MUTED)
        add_text(s, Inches(0.7), Inches(2.55), Inches(11), Inches(0.3), pr["meta"], 11, False, FAINT)

        # three image frames
        xs = [Inches(0.7), Inches(4.85), Inches(9.0)]
        for j, rel in enumerate(pr["imgs"]):
            p = heji / rel if not rel.startswith("assets/") else ASSETS / rel.replace("assets/", "")
            if not p.exists():
                p = ASSETS / Path(rel).name
            add_cover_image(s, p, xs[j], Inches(3.05), Inches(3.4), Inches(3.55))
            add_rect(s, xs[j], Inches(3.05), Inches(3.4), Inches(3.55), None, INK, 2)
        footer(s, 4 + i)
        notes(s, pr["title"] + " 项目展示。")

    # ========== 7 AI ==========
    s = prs.slides.add_slide(blank)
    add_bg(s)
    section_label(s, Inches(0.7), Inches(0.45), "AI 视频影像", "AI SHOWREEL")
    add_text(s, Inches(0.7), Inches(0.95), Inches(11), Inches(0.5), "生成影像实验场", 28, True, INK)
    add_text(s, Inches(0.7), Inches(1.55), Inches(11), Inches(0.35),
             "ComfyUI 串联 MiniMax H3 与 Seedance 2.0，I2V / T2V 影像。正在学习积累 · 尚未完善", 13, False, MUTED)
    add_cover_image(s, ASSETS / "ai-1.png", Inches(0.7), Inches(2.2), Inches(5.8), Inches(4.2))
    add_rect(s, Inches(0.7), Inches(2.2), Inches(5.8), Inches(4.2), None, INK, 2)
    add_cover_image(s, ASSETS / "ai-2.png", Inches(6.85), Inches(2.2), Inches(5.8), Inches(4.2))
    add_rect(s, Inches(6.85), Inches(2.2), Inches(5.8), Inches(4.2), None, INK, 2)
    footer(s, 7)
    notes(s, "AI 视频影像。")

    # ========== 8 Interactive ==========
    s = prs.slides.add_slide(blank)
    add_bg(s)
    section_label(s, Inches(0.7), Inches(0.45), "交互与数字孪生", "INTERACTIVE / TWIN")
    add_text(s, Inches(0.7), Inches(0.95), Inches(11), Inches(0.5), "可上手把玩的实验", 28, True, INK)

    cards = [
        ("交互网站", "像素画板 · 全屋工坊", "正在学习积累 · 尚未完善", ASSETS / "xiang-1.png"),
        ("室内设计工坊", "灵感家 · 3D 全屋", "材质 / 家具 / 光照", ASSETS / "fang-2.png"),
        ("数字孪生", "云璟天悦智慧社区", "Three.js + 数据看板", ASSETS / "twin-2.png"),
    ]
    for j, (tag, title, sub, img) in enumerate(cards):
        x = Inches(0.7 + j * 4.15)
        add_cover_image(s, img, x, Inches(1.7), Inches(3.9), Inches(2.6))
        add_rect(s, x, Inches(1.7), Inches(3.9), Inches(2.6), None, INK, 2)
        title_bar(s, x, Inches(4.5), Inches(3.9), Inches(1.55), title, tag.upper(), sub)
    footer(s, 8)
    notes(s, "交互与孪生项目。")

    # ========== 9 Projects list ==========
    s = prs.slides.add_slide(blank)
    add_bg(s)
    section_label(s, Inches(0.7), Inches(0.45), "项目索引", "PROJECTS")
    add_text(s, Inches(0.7), Inches(0.95), Inches(11), Inches(0.45), "全部项目", 26, True, INK)

    rows = [
        ("01", "青峦望海 · 青山村", "乡村景观更新 · 青岛崂山"),
        ("02", "泉汇西枢 · 济南西中央公园", "城市中央公园 · 济南西站"),
        ("03", "流·筑未来 · 企业共享花园", "企业园区景观 · 淄博"),
        ("04", "AI 视频专区", "MiniMax H3 / Seedance 2.0"),
        ("05", "交互网站设计", "3D 全屋工坊 · 像素画板"),
        ("06", "数字孪生 · 云璟天悦", "智慧社区数据看板"),
        ("07", "回澜阁 · 线上展示站", "wangfuxin-design.github.io"),
        ("08", "像素画板 / 室内工坊", "纯前端交互实验"),
    ]
    y = Inches(1.55)
    for i, (n, t, sub) in enumerate(rows):
        list_row(s, Inches(0.7), y, Inches(12.0), n, t, sub, accent=(i % 2 == 1))
        y += Inches(0.68)
    footer(s, 9)
    notes(s, "项目索引列表。")

    # ========== 10 Contact ==========
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_rect(s, 0, 0, Inches(0.12), SH, ORANGE, None)
    add_text(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(0.5), "一起做点东西", 36, True, INK)
    orange_bar(s, Inches(0.92), Inches(2.5), Inches(1.2))
    add_text(s, Inches(0.9), Inches(2.85), Inches(10), Inches(0.8),
             "空间、设计，与一切有意思的事，我都感兴趣。", 22, True, ORANGE)

    boxes = [
        ("EMAIL", "wx2903964065@126.com"),
        ("PORTFOLIO", "wangfuxin-design.github.io"),
        ("SCHOOL", "青岛理工大学 · 环境设计"),
    ]
    for i, (k, v) in enumerate(boxes):
        x = Inches(0.9 + i * 4.0)
        add_rect(s, x, Inches(4.2), Inches(3.7), Inches(1.5), WHITE, INK, 2)
        add_rect(s, x + Inches(0.05), Inches(4.25), Inches(3.7), Inches(1.5), None, INK, 2)
        add_text(s, x + Inches(0.28), Inches(4.45), Inches(3.2), Inches(0.3), k, 11, True, ORANGE, font=FONT_EN)
        add_text(s, x + Inches(0.28), Inches(4.9), Inches(3.2), Inches(0.5), v, 14, True, INK)

    add_text(s, Inches(0.9), Inches(6.2), Inches(10), Inches(0.35),
             "有任何合作意向，请通过邮箱联系我。", 12, False, MUTED)
    footer(s, 10)
    notes(s, "联系方式。")

    # ========== 11 End ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, INK)
    add_text(s, Inches(0), Inches(2.8), SW, Inches(0.8), "方寸之间 · 想象之上", 36, True, WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(0), Inches(3.7), SW, Inches(0.4), "WANG FUXIN  ·  ENVIRONMENTAL DESIGN  ·  2026", 12, False, RGBColor(0xAA, 0xAA, 0xAA), PP_ALIGN.CENTER, FONT_EN)
    add_rect(s, Inches(6.1), Inches(4.4), Inches(1.1), Emu(28575), ORANGE, None)
    add_text(s, Inches(0), Inches(4.8), SW, Inches(0.4), "Thanks", 16, False, ORANGE, PP_ALIGN.CENTER, FONT_EN)
    notes(s, "谢谢观看。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print("saved", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    build()
